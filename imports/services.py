from openpyxl import load_workbook

from books.models import LibraryItem
from .utils import get_value


def analyze_excel(file):

    workbook = load_workbook(file)

    sheet = workbook.active

    rows = list(sheet.iter_rows(values_only=True))

    if len(rows) < 2:
        return None

    headers = [
        str(h).strip() if h else ""
        for h in rows[0]
    ]

    header_map = {
        header: index
        for index, header in enumerate(headers)
    }

    results = []

    total = 0
    new_books = 0
    existing_books = 0
    duplicate_rows = 0
    invalid_rows = 0

    seen = set()

    for row in rows[1:]:

        # Ignore completely empty rows
        if all(
            cell is None or str(cell).strip() == ""
            for cell in row
        ):
            continue

        total += 1

        row_data = {}

        for header, index in header_map.items():

            row_data[header] = (
                row[index]
                if index < len(row)
                else ""
            )

        accession = str(
            get_value(
                row_data,
                "accession_number",
            ) or ""
        ).strip()

        title = str(
            get_value(
                row_data,
                "title",
            ) or ""
        ).strip()

        author = str(
            get_value(
                row_data,
                "author",
            ) or ""
        ).strip()

        publisher = str(
            get_value(
                row_data,
                "publisher",
            ) or ""
        ).strip()

        # Skip invalid rows
        if not accession or not title:

            invalid_rows += 1
            continue

        # Duplicate inside Excel
        if accession in seen:

            duplicate_rows += 1

            results.append({

                "accession_number": accession,
                "title": title,
                "author": author,
                "publisher": publisher,
                "status": "Duplicate",

            })

            continue

        seen.add(accession)

        exists = LibraryItem.objects.filter(
            accession_number=accession
        ).exists()

        if exists:

            status = "Exists"
            existing_books += 1

        else:

            status = "New"
            new_books += 1

        results.append({

            "accession_number": accession,
            "title": title,
            "author": author,
            "publisher": publisher,
            "status": status,

        })

    return {

        "headers": headers,
        "results": results,
        "total": total,
        "new_books": new_books,
        "existing_books": existing_books,
        "duplicate_rows": duplicate_rows,
        "invalid_rows": invalid_rows,

    }
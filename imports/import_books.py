from openpyxl import load_workbook
import re

from donors.models import Donor
from .constants import COLUMN_ALIASES
from .utils import get_value

from books.models import (
    Author,
    Category,
    LibraryItem,
    Publisher,
)


def import_books(excel_file):

    workbook = load_workbook(excel_file)
    sheet = workbook.active

    rows = list(sheet.iter_rows(values_only=True))

    if not rows:
        return {
            "imported": 0,
            "skipped": 0,
            "duplicates": 0,
        }

    headers = [
        str(h).strip() if h else ""
        for h in rows[0]
    ]

    print(headers)

    header_map = {
        header: index
        for index, header in enumerate(headers)
    }
    print(headers)
    print(header_map.keys())

    general_category, _ = Category.objects.get_or_create(
        name="General"
    )

    seen = set()

    imported = 0
    skipped = 0
    duplicates = 0

    for row in rows[1:]:

        # Skip completely empty rows
        if all(
            cell is None or str(cell).strip() == ""
            for cell in row
        ):
            continue

        row_data = {}

        for header, index in header_map.items():

            row_data[header] = (
                row[index]
                if index < len(row)
                else ""
            )

        accession = str(
            get_value(
                row_data,
                "accession_number",
            ) or ""
        ).strip()

        title = str(
            get_value(
                row_data,
                "title",
            ) or ""
        ).strip()

        if not accession or not title:
            continue

        if accession in seen:
            duplicates += 1
            continue

        seen.add(accession)

        if LibraryItem.objects.filter(
            accession_number=accession
        ).exists():

            skipped += 1
            continue

        author_name = str(
            get_value(
                row_data,
                "author",
            ) or "Unknown"
        ).strip()

        publisher_name = str(
            get_value(
                row_data,
                "publisher",
            ) or "Unknown"
        ).strip()

        author, _ = Author.objects.get_or_create(
            name=author_name
        )

        publisher, _ = Publisher.objects.get_or_create(
            name=publisher_name
        )

        source = str(
            get_value(
                row_data,
                "source",
            ) or ""
        ).strip()

        print(repr(source))

        donor = None
        acquisition_type = LibraryItem.PURCHASE

        if source.lower().startswith("donated"):

            acquisition_type = LibraryItem.DONATION

            text = source.replace("\n", " ")

            employee = ""

            match = re.search(
                r"EC\s*No\.?\s*([0-9]+)",
                text,
                re.IGNORECASE,
            )

            if match:
                employee = match.group(1)

            donor_name = re.sub(
                r"(?i)donated\s+by",
                "",
                text,
            )

            donor_name = re.sub(
                r"EC\s*No\.?\s*[0-9]+",
                "",
                donor_name,
            )

            donor_name = donor_name.strip(" ,")

            donor, created = Donor.objects.get_or_create(
                name=donor_name,
                defaults={
                    "employee_code": employee,
                },
            )

            if employee and not donor.employee_code:
                donor.employee_code = employee
                donor.save()

        category_name = str(
            get_value(
                row_data,
                "category",
            ) or ""
        ).strip()

        if category_name:

            category, _ = Category.objects.get_or_create(
                name=category_name,
            )

        else:

            category = general_category

        print("ROW DATA KEYS:")
        for key in row_data.keys():
            print(repr(key))

        print("PRICE ALIASES:", COLUMN_ALIASES["price"])


        print("PRICE FROM EXCEL:", repr(get_value(row_data, "price")))        

        LibraryItem.objects.create(

            accession_number=accession,

            date_added=get_value(
                row_data,
                "date_added",
            ),

            serial_number=get_value(
                row_data,
                "serial_number",
            ),

            title=title,

            author=author,

            publisher=publisher,

            category=category,

            publication_year=get_value(
                row_data,
                "publication_year",
            ),

            pages=get_value(
                row_data,
                "pages",
            ),

            price=row_data["Rate in Rs."],

            volume_qty=(
                get_value(
                    row_data,
                    "volume_qty",
                )
                or 1
            ),

            source=source,

            acquisition_type=acquisition_type,

            donor=donor,

            location=get_value(
                row_data,
                "location",
            ) or "",

            remarks=get_value(
                row_data,
                "remarks",
            ) or "",

            status=LibraryItem.AVAILABLE,
        )

        imported += 1

    return {
        "imported": imported,
        "skipped": skipped,
        "duplicates": duplicates,
    }